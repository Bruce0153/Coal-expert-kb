from __future__ import annotations

from typing import List

from langchain_core.documents import Document

from coal_kb.loaders.base import BaseLoader, detect_language, normalize_text
from coal_kb.loaders.registry import register_loader


class XlsxLoader(BaseLoader):
    def __init__(self) -> None:
        try:
            import openpyxl  # noqa: F401
        except Exception as exc:  # pragma: no cover - optional dependency
            raise RuntimeError("openpyxl is required for XlsxLoader") from exc

    def load(self, path: str) -> List[Document]:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        docs: List[Document] = []
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            chunk = []
            start = 1
            for idx, row in enumerate(rows, start=1):
                chunk.append(", ".join("" if v is None else str(v) for v in row))
                if len(chunk) >= 50:
                    content = normalize_text("\n".join(chunk))
                    lang = detect_language(content)
                    docs.append(
                        Document(
                            page_content=content,
                            metadata={
                                "source_file": path,
                                "section": sheet.title,
                                "row_range": [start, idx],
                                "doc_type": "xlsx",
                                "language": lang,
                                "parser": "xlsx",
                            },
                        )
                    )
                    chunk = []
                    start = idx + 1
            if chunk:
                content = normalize_text("\n".join(chunk))
                lang = detect_language(content)
                docs.append(
                    Document(
                        page_content=content,
                        metadata={
                            "source_file": path,
                            "section": sheet.title,
                            "row_range": [start, len(rows)],
                            "doc_type": "xlsx",
                            "language": lang,
                            "parser": "xlsx",
                        },
                    )
                )
        return docs


register_loader("xlsx", XlsxLoader)
